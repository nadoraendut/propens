"""The draw: select auditees and auditors fairly, explain the result, and build the Excel workbook."""
import math
import random
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from propens import logic, rotation

LEDGER_HEADERS = ["Semester", "Staff ID", "Staff Name", "Bidang", "In Pool", "Auditee", "Auditor Count"]


def _lecturer(l):
    return dict(l, id=l["staff_id"], start_date=str(l["start_date"]))


def run_draw(lecturers, ledger, settings, today=None):
    """lecturers: dicts (staff_id, name, grade, bidang, start_date, campus, status).
    ledger: canonical history rows. settings: semester, percent, seed, exclude, force."""
    today = today or date.today()
    name = settings["semester"].strip()
    percent = float(settings["percent"])
    seed = settings["seed"]
    rng = random.Random(seed)
    warnings = []

    everyone = [_lecturer(l) for l in lecturers]
    pool = [p for p in everyone if logic.is_eligible(p["status"])]
    if not pool:
        raise ValueError("No lecturer with Status 'Active' in the list.")
    inactive = len(everyone) - len(pool)

    same = [r for r in ledger if r["semester"] == name]
    if same:                                  # drawing the same semester again: it must not count against itself
        ledger = [r for r in ledger if r["semester"] != name]
        warnings.append(f"The history already had '{name}'. It was replaced by this draw.")
    history = rotation.parse_ledger(ledger)

    known = {p["staff_id"] for p in pool}
    exclude, force = set(settings["exclude"]), set(settings["force"])
    for label, ids in (("exclude", exclude), ("must include", force)):
        for i in sorted(ids - known):
            warnings.append(f"Staff ID {i} ({label}) is not an active lecturer in the list and was ignored.")
    exclude &= known
    force &= known
    clash = exclude & force
    if clash:
        raise ValueError("Staff ID(s) both excluded and forced: " + ", ".join(sorted(clash)))

    by = defaultdict(list)
    for p in pool:
        by[p["bidang"]].append(p)
    st = rotation.standing(pool, history, today)
    chosen, short = rotation.select_auditees(by, percent, st, rng, exclude, force)
    pairing = rotation.assign_auditors(chosen, pool, st, rng, today)
    byid = {p["id"]: p for p in pool}

    rows = []
    for a in sorted(chosen, key=lambda x: (x["bidang"], x["name"])):
        er = byid.get(pairing[a["id"]])
        if er is None:
            rule = "none"
        else:
            rule = "ok" if logic.can_audit(er, a, today) else "peer"
        rows.append({
            "staff_id": a["staff_id"], "name": a["name"], "grade": a["grade"], "bidang": a["bidang"],
            "campus": a.get("campus", ""), "years": logic.years_of_service(a["start_date"], today),
            "waiting": st[a["id"]]["waiting"], "reason": rotation.reason(st[a["id"]]),
            "forced": a["id"] in force,
            "auditor_staff_id": er["staff_id"] if er else "", "auditor_name": er["name"] if er else "",
            "auditor_grade": er["grade"] if er else "",
            "auditor_years": logic.years_of_service(er["start_date"], today) if er else None,
            "rule": rule})

    selected_by_dept = defaultdict(int)
    for r in rows:
        selected_by_dept[r["bidang"]] += 1
    departments = []
    for d in sorted(by):
        need = logic.required_count(len(by[d]), percent)
        departments.append({"dept": d, "total": len(by[d]), "required": need, "selected": selected_by_dept[d],
                            "ok": selected_by_dept[d] >= need,
                            "cycle": math.ceil(len(by[d]) / need) if need else None})

    chosen_ids = {a["id"] for a in chosen}
    prev_name = history[-1]["name"] if history else None
    repeats = [r for r in rows if history and r["waiting"] == 0 and st[r["staff_id"]]["times"] > 0
               and r["staff_id"] in history[-1]["audited"]]
    pool_wait = [st[p["id"]]["waiting"] for p in pool]
    sel_wait = [r["waiting"] for r in rows]
    never_left = sum(1 for p in pool if st[p["id"]]["times"] == 0 and p["id"] not in chosen_ids)
    fairness = {
        "history_semesters": len(history), "previous": prev_name,
        "avg_wait_selected": round(sum(sel_wait) / len(sel_wait), 2) if sel_wait else None,
        "avg_wait_pool": round(sum(pool_wait) / len(pool_wait), 2) if pool_wait else None,
        "repeats": [{"staff_id": r["staff_id"], "name": r["name"], "bidang": r["bidang"]} for r in repeats],
        "never_audited_remaining": never_left if history else None,
        "audited_at_least_once": sum(1 for p in pool if st[p["id"]]["times"] > 0 or p["id"] in chosen_ids),
        "pool": len(pool),
    }

    served = defaultdict(int)
    for v in pairing.values():
        if v:
            served[v] += 1
    workload = sorted(({"staff_id": byid[i]["staff_id"], "name": byid[i]["name"], "grade": byid[i]["grade"],
                        "assigned": n, "past_total": st[i]["auditor_total"],
                        "also_audited": i in chosen_ids} for i, n in served.items()),
                      key=lambda w: (-w["assigned"], w["name"]))
    new_ledger = rotation.ledger_rows(name, pool, chosen_ids, pairing, None)
    carried = [{"semester": r["semester"], "staff_id": r["staff_id"], "name": r.get("name") or "",
                "bidang": r.get("bidang") or "",
                "in_pool": r.get("in_pool") or "",
                "auditee": r["auditee"] if r.get("auditee") not in (None, "") else "Yes",
                "auditor_count": r.get("auditor_count") or 0} for r in ledger]
    return {
        "semester": name, "percent": percent, "seed": seed, "date": today.isoformat(),
        "rows": rows, "departments": departments,
        "shortfalls": [{"dept": d, "needed": n, "got": g} for d, n, g in short],
        "unmatched": sum(1 for r in rows if r["rule"] == "none"),
        "peers": sum(1 for r in rows if r["rule"] == "peer"),
        "inactive": inactive, "fairness": fairness, "workload": workload, "warnings": warnings,
        "ledger": carried + new_ledger,
    }


def build_workbook(result):
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit List"
    ws.append(["Semester", "Staff ID", "Staff Name", "Gred Jawatan", "Bidang", "Campus", "Years of Service",
               "Auditor Staff ID", "Auditor Name", "Auditor Gred", "Auditor Years", "Auditor Rule",
               "Why selected"])
    rules = {"ok": "Higher grade / more senior", "peer": "Same seniority (no one outranks)", "none": "HEAD TO PICK"}
    for r in result["rows"]:
        ws.append([result["semester"], r["staff_id"], r["name"], r["grade"], r["bidang"], r["campus"], r["years"],
                   r["auditor_staff_id"], r["auditor_name"], r["auditor_grade"], r["auditor_years"],
                   rules[r["rule"]], r["reason"] + (" (must include)" if r["forced"] else "")])
    d = wb.create_sheet("Departments")
    d.append(["Bidang", "Active lecturers", f"Required (min {result['percent']:g}%)", "Selected", "OK",
              "Semesters for everyone to have a turn"])
    for x in result["departments"]:
        d.append([x["dept"], x["total"], x["required"], x["selected"], "Yes" if x["ok"] else "NO", x["cycle"]])
    f = wb.create_sheet("Fairness")
    fr = result["fairness"]
    f.append(["Item", "Value"])
    f.append(["Semester", result["semester"]])
    f.append(["Drawn on", result["date"]])
    f.append(["Random seed (same files + settings + seed = same list)", result["seed"]])
    f.append(["Semesters of history used", fr["history_semesters"]])
    f.append(["Average semesters waited: selected lecturers", fr["avg_wait_selected"]])
    f.append(["Average semesters waited: all active lecturers", fr["avg_wait_pool"]])
    f.append(["Selected again straight after the previous semester", len(fr["repeats"])])
    f.append(["Lecturers still never audited after this draw", fr["never_audited_remaining"]])
    f.append(["Lecturers audited at least once (incl. this draw)", f"{fr['audited_at_least_once']} of {fr['pool']}"])
    for x in fr["repeats"]:
        f.append([f"  Repeat: {x['name']} ({x['staff_id']})", x["bidang"] + " - department too small to avoid it"])
    w = wb.create_sheet("Auditor Workload")
    w.append(["Staff ID", "Auditor", "Gred", "Audits this semester", "Audits in history", "Also being audited"])
    for x in result["workload"]:
        w.append([x["staff_id"], x["name"], x["grade"], x["assigned"], x["past_total"],
                  "Yes" if x["also_audited"] else "No"])
    lg = wb.create_sheet("Ledger")
    lg.append(LEDGER_HEADERS)
    for r in result["ledger"]:
        lg.append([r["semester"], r["staff_id"], r.get("name", ""), r.get("bidang", ""), r.get("in_pool", ""),
                   r.get("auditee", ""), r.get("auditor_count", 0)])
    fill = PatternFill("solid", fgColor="DDE6F3")
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 55)
        sheet.freeze_panes = "A2"
    return wb


def ledger_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(LEDGER_HEADERS)
    ws.append(["Sem 1 2025/2026", "A1001", "Dr. Example", "Computer Science", "Yes", "Yes", 0])
    ws.append(["Sem 1 2025/2026", "A1002", "Dr. Another", "Computer Science", "Yes", "No", 1])
    n = wb.create_sheet("How to fill")
    for line in ["Oldest semester first. One row per lecturer per semester.",
                 "Quick start if you only know who was audited: keep just Semester + Staff ID (one row per audited "
                 "lecturer). Leave the other columns out.",
                 "Full format (what this tool exports): In Pool = Yes/No, Auditee = Yes/No, Auditor Count = how many "
                 "lecturers they audited.",
                 "Each semester, upload the workbook this tool gave you last time: its 'Ledger' sheet already contains "
                 "all earlier semesters."]:
        n.append([line])
    n.column_dimensions["A"].width = 110
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFG", (22, 12, 24, 24, 10, 10, 14)):
        ws.column_dimensions[col].width = w
    return wb
