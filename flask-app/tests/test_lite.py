import io
import json
import re
from html import unescape

import pytest
from openpyxl import load_workbook

from lite.app import app as lite_app


def lecturers_csv(extra=""):
    rows = ["Staff ID,Staff Name,Gred Jawatan,Bidang,Start of Employment,Campus,Status,Email"]
    n = 0
    for dept, size in [("CS", 10), ("Math", 5), ("Physics", 4)]:
        for i in range(size):
            n += 1
            rows.append(f"S{n:03d},Lect {n},DS{41 + (n % 4) * 3},{dept},{2000 + n}-01-01,N,Active,s{n}@u.edu")
    rows.append("S900,On Leave,DS52,CS,2001-01-01,N,Study Leave,x@u.edu")
    return ("\n".join(rows) + extra + "\n").encode()


@pytest.fixture
def client(monkeypatch):
    monkeypatch.delenv("LITE_PASSWORD", raising=False)
    return lite_app.test_client()


def draw(client, history=None, history_name="history.xlsx", **form):
    data = {"semester": "Sem A", "percent": "20", "lecturers": (io.BytesIO(lecturers_csv()), "l.csv")}
    if history is not None:
        data["history"] = (io.BytesIO(history), history_name)
    data.update(form)
    return client.post("/generate", data=data, content_type="multipart/form-data")


def hidden(html, name):
    return unescape(re.search(rf'name="{name}" value="([^"]*)"', html).group(1))


def download(client, html):
    r = client.post("/download", data={"result": hidden(html, "result")})
    assert r.status_code == 200
    return r.data


def selected_ids(xlsx):
    ws = load_workbook(io.BytesIO(xlsx))["Audit List"]
    return {r[1] for r in ws.iter_rows(min_row=2, values_only=True)}


def test_first_draw_meets_quota_and_rule(client):
    r = draw(client, seed="42")
    assert r.status_code == 200
    html = r.data.decode()
    assert "Seed <b>42</b>" in html
    wb = load_workbook(io.BytesIO(download(client, html)))
    everyone = {c for row in wb["Audit List"].iter_rows(min_row=2, values_only=True) for c in (row[1], row[7])}
    assert "S900" not in everyone                                          # Study Leave: not drawn, not an auditor
    assert {"Audit List", "Departments", "Fairness", "Auditor Workload", "Ledger"} <= set(wb.sheetnames)
    rows = list(wb["Audit List"].iter_rows(min_row=2, values_only=True))
    by_dept = {}
    for row in rows:
        by_dept[row[4]] = by_dept.get(row[4], 0) + 1
    assert by_dept == {"CS": 2, "Math": 1, "Physics": 1}                  # ceil(20%) per Bidang
    assert all(row[7] and row[7] != row[1] for row in rows)               # everyone has an auditor, never themselves
    ledger = list(wb["Ledger"].iter_rows(min_row=2, values_only=True))
    assert len(ledger) == 19 and sum(1 for x in ledger if x[5] == "Yes") == 4   # full active pool; 4 audited


def test_same_seed_reproduces_the_list(client):
    a = selected_ids(download(client, draw(client, seed="7").data.decode()))
    b = selected_ids(download(client, draw(client, seed="7").data.decode()))
    assert a == b


def test_downloaded_workbook_rotates_next_semester(client):
    history, picked = None, []
    for n in range(1, 6):                       # CS has 10 active lecturers, quota 2 -> a 5-semester cycle
        html = draw(client, history, semester=f"Sem {n}").data.decode()
        xlsx = download(client, html)
        picked.append({i for i in selected_ids(xlsx) if int(i[1:]) <= 10})
        history = xlsx                          # exactly what the committee would do: upload last semester's file
    assert all(len(p) == 2 for p in picked)
    assert len(set().union(*picked)) == 10 and not any(a & b for i, a in enumerate(picked) for b in picked[i + 1:])
    html = draw(client, history, semester="Sem 6").data.decode()
    assert len(load_workbook(io.BytesIO(download(client, html)))["Ledger"]["A"]) == 6 * 19 + 1   # history carried forward


def test_simple_history_list_is_accepted(client):
    simple = "Semester,Staff ID\nSem 0,S001\nSem 0,S002\nSem 0,S003\n".encode()
    r = client.post("/generate", data={"semester": "Sem 1", "percent": "20", "seed": "1",
                                       "lecturers": (io.BytesIO(lecturers_csv()), "l.csv"),
                                       "history": (io.BytesIO(simple), "h.csv")}, content_type="multipart/form-data")
    ids = selected_ids(download(client, r.data.decode()))
    assert not ids & {"S001", "S002", "S003"}                              # audited last semester, others are overdue


def test_draw_again_and_exclude_force(client):
    html = draw(client, seed="3", exclude="S001 S002", force="S010").data.decode()
    ids = selected_ids(download(client, html))
    assert "S010" in ids and not ids & {"S001", "S002"} and "must include" in html
    again = client.post("/generate", data={"payload": hidden(html, "payload"), "semester": "Sem A", "percent": "20",
                                           "exclude": "S001 S002", "force": "S010"})        # no re-upload needed
    assert again.status_code == 200 and "S010" in selected_ids(download(client, again.data.decode()))


def test_bad_inputs_are_explained_not_crashed(client):
    r = client.post("/generate", data={"semester": "X", "percent": "20"})
    assert r.status_code == 400 and "Upload the lecturer list" in r.data.decode()
    r = draw(client, percent="150")
    assert r.status_code == 400 and "between 0 and 100" in r.data.decode()
    bad = lecturers_csv(extra="\nS901,Bad,DS52,CS,notadate,N,Active,a@u.edu")
    r = client.post("/generate", data={"semester": "X", "percent": "20", "lecturers": (io.BytesIO(bad), "l.csv")},
                    content_type="multipart/form-data")
    assert r.status_code == 400 and "Row" in r.data.decode() and "nothing was drawn" in r.data.decode()
    r = draw(client, history=b"Foo,Bar\n1,2\n", history_name="h.csv")
    assert r.status_code == 400 and "Semester, Staff ID" in r.data.decode()
    r = draw(client, history=b"Semester,Staff ID\nA,1\n", history_name="h.xlsx")           # a CSV pretending to be Excel
    assert r.status_code == 400 and "could not read the file as an Excel" in r.data.decode()
    r = draw(client, exclude="S001", force="S001")
    assert r.status_code == 400 and "both excluded and forced" in r.data.decode()
    assert client.post("/download", data={"result": "not json"}).status_code == 400


def test_redrawing_same_semester_replaces_it_in_history(client):
    first = download(client, draw(client, semester="Sem A").data.decode())
    html = draw(client, first, semester="Sem A").data.decode()
    assert "replaced by this draw" in html
    ledger = list(load_workbook(io.BytesIO(download(client, html)))["Ledger"].iter_rows(min_row=2, values_only=True))
    assert len(ledger) == 19                                                  # not duplicated


def test_password_gate(monkeypatch):
    monkeypatch.setenv("LITE_PASSWORD", "s3cret")
    c = lite_app.test_client()
    assert c.get("/").status_code == 401 and c.post("/generate").status_code == 401
    import base64
    ok = {"Authorization": "Basic " + base64.b64encode(b"me:s3cret").decode()}
    bad = {"Authorization": "Basic " + base64.b64encode(b"me:wrong").decode()}
    assert c.get("/", headers=ok).status_code == 200 and c.get("/", headers=bad).status_code == 401


def test_vercel_entrypoint_exposes_the_same_app(client):
    from api.index import app as vercel_app
    assert vercel_app is lite_app and client.get("/").status_code == 200
    assert client.get("/lecturer-template.xlsx").status_code == 200 and client.get("/history-template.xlsx").status_code == 200
