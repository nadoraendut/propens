"""Workflow operations on a semester's audit list."""
import csv
import io
import random
import re
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.security import generate_password_hash

from . import logic, rotation, rubric
from .db import get_db
from .notify import now, notify, send_email, tell, users_with_role

EDITABLE_STATES = ("draft", "generated", "changes_requested")

COLUMNS = [  # (key, header, accepted header aliases)
    ("staff_id", "Staff ID", ["staffid", "id", "staffno"]),
    ("name", "Staff Name", ["staffname", "name", "nama"]),
    ("grade", "Gred Jawatan", ["gredjawatan", "grade", "gred"]),
    ("bidang", "Bidang", ["bidang", "department", "field"]),
    ("start_date", "Start of Employment", ["startofemployment", "startdate", "tarikhlantikan"]),
    ("campus", "Campus", ["campus", "kampus"]),
    ("status", "Status", ["status"]),
    ("email", "Email", ["email", "emel", "emailaddress"]),
]
REQUIRED = ("staff_id", "name", "grade", "bidang", "start_date")

# Mark bands used in reports (TAS out of 100) -- the indicator table printed on the PRO-PENS form.
BANDS = [(f"{label} ({'<60' if floor == 0 else (str(floor) + '-' + ('100' if floor == 90 else str(floor + 9)))})", floor)
         for floor, label in rubric.INDICATORS]


# ---------------------------------------------------------------- importing
def _norm(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def read_rows(filename, data):
    """Return list of header-keyed dict rows from an .xlsx or .csv upload."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            ws = load_workbook(io.BytesIO(data), data_only=True).active
        except Exception:
            raise ValueError("Could not read the file as an Excel (.xlsx) workbook. Is it damaged, or a CSV "
                             "saved with an .xlsx name?") from None
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    elif filename.lower().endswith(".csv"):
        raw = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        raise ValueError("Upload an .xlsx or .csv file.")
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if not raw:
        raise ValueError("The file is empty.")
    alias = {a: key for key, _h, als in COLUMNS for a in als + [_norm(_h)]}
    cols = {i: alias[_norm(h)] for i, h in enumerate(raw[0]) if _norm(h) in alias}
    missing = [h for k, h, _ in COLUMNS if k in REQUIRED and k not in cols.values()]
    if missing:
        raise ValueError("Missing column(s): " + ", ".join(missing))
    return [{key: (r[i] if i < len(r) else None) for i, key in cols.items()} for r in raw[1:]]


def clean_lecturer_rows(rows):
    """Validate uploaded rows. Returns (good rows, error strings); bad rows are skipped, good ones kept."""
    errors, good, seen = [], [], set()
    headers = dict((c[0], c[1]) for c in COLUMNS)
    for n, row in enumerate(rows, start=2):
        try:
            clean = {k: ("" if row.get(k) is None else str(row.get(k)).strip())
                     for k in ("staff_id", "name", "grade", "bidang", "campus", "status", "email")}
            clean["status"] = clean["status"] or "Active"
            for k in ("staff_id", "name", "grade", "bidang"):
                if not clean[k]:
                    raise ValueError(f"{headers[k]} is empty")
            if clean["staff_id"] in seen:
                raise ValueError("duplicate Staff ID in file")
            seen.add(clean["staff_id"])
            clean["start_date"] = logic.parse_date(row.get("start_date"))
            if clean["start_date"] is None:
                raise ValueError("Start of Employment is empty")
            if clean["start_date"] > date.today():
                raise ValueError("Start of Employment is in the future")
            good.append(clean)
        except ValueError as exc:
            errors.append(f"Row {n}: {exc}")
    return good, errors


def import_lecturers(semester_id, rows, replace=False):
    """Upsert lecturers by Staff ID and link them to the semester.

    Returns (imported_count, [error strings]). Bad rows are skipped, good rows kept.
    """
    db = get_db()
    good, errors = clean_lecturer_rows(rows)
    if replace:
        db.execute("DELETE FROM semester_lecturers WHERE semester_id=?", (semester_id,))
    for c in good:
        db.execute(
            """INSERT INTO lecturers (staff_id,name,grade,bidang,start_date,campus,status,email,password_hash,
                                      must_change_password)
               VALUES (?,?,?,?,?,?,?,?,?,1)
               ON CONFLICT(staff_id) DO UPDATE SET name=excluded.name, grade=excluded.grade,
                 bidang=excluded.bidang, start_date=excluded.start_date, campus=excluded.campus,
                 status=excluded.status, email=CASE WHEN excluded.email<>'' THEN excluded.email ELSE email END""",
            (c["staff_id"], c["name"], c["grade"], c["bidang"], c["start_date"].isoformat(),
             c["campus"], c["status"], c["email"], generate_password_hash(c["staff_id"])))
        lid = db.execute("SELECT id FROM lecturers WHERE staff_id=?", (c["staff_id"],)).fetchone()["id"]
        db.execute("INSERT OR IGNORE INTO semester_lecturers VALUES (?,?)", (semester_id, lid))
    db.commit()
    return len(good), errors


def template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Lecturers"
    ws.append([h for _k, h, _a in COLUMNS])
    ws.append(["A1001", "Dr. Example Name", "DS52", "Computer Science", "2012-03-01",
               "Main Campus", "Active", "example@university.edu"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 22
    return wb


# ---------------------------------------------------------------- queries
def semester_lecturers(semester_id):
    return get_db().execute(
        """SELECT l.* FROM lecturers l JOIN semester_lecturers sl ON sl.lecturer_id=l.id
           WHERE sl.semester_id=? ORDER BY l.bidang, l.name""", (semester_id,)).fetchall()


def audit_rows(semester_id):
    return get_db().execute(
        """SELECT a.*, ee.staff_id ee_staff, ee.name ee_name, ee.bidang ee_bidang, ee.campus ee_campus,
                  ee.grade ee_grade, ee.start_date ee_start, ee.email ee_email,
                  er.staff_id er_staff, er.name er_name, er.grade er_grade, er.start_date er_start,
                  er.email er_email
           FROM audits a JOIN lecturers ee ON ee.id=a.auditee_id
           LEFT JOIN lecturers er ON er.id=a.auditor_id
           WHERE a.semester_id=? ORDER BY ee.bidang, ee.name""", (semester_id,)).fetchall()


def rule_statuses(semester_id):
    """{audit_id: 'ok' | 'peer' | 'override' | 'none'} for every audit in the semester.

    ok       auditor outranks the auditee (grade) or is more senior (years)
    peer     nobody outranks the auditee, and the auditor has the same years of service
    override any other manual pick
    none     no auditor yet
    """
    pool = [dict(l) for l in semester_lecturers(semester_id) if logic.is_eligible(l["status"])]
    out = {}
    for a in audit_rows(semester_id):
        if a["auditor_id"] is None:
            out[a["id"]] = "none"
            continue
        ee = {"id": a["auditee_id"], "grade": a["ee_grade"], "start_date": a["ee_start"]}
        er = {"id": a["auditor_id"], "grade": a["er_grade"], "start_date": a["er_start"]}
        if logic.can_audit(er, ee):
            out[a["id"]] = "ok"
        elif logic.is_peer(er, ee) and not any(logic.can_audit(p, ee) for p in pool):
            out[a["id"]] = "peer"
        else:
            out[a["id"]] = "override"
    return out


def department_quota(semester):
    """Per-department required vs selected counts."""
    people = [l for l in semester_lecturers(semester["id"]) if logic.is_eligible(l["status"])]
    by = defaultdict(list)
    for p in people:
        by[p["bidang"]].append(p)
    audits = audit_rows(semester["id"])
    out = []
    for dept in sorted(by):
        sel = [a for a in audits if a["ee_bidang"] == dept and a["round_no"] == 1]
        need = logic.required_count(len(by[dept]), semester["min_percent"])
        out.append({"dept": dept, "total": len(by[dept]), "required": need,
                    "selected": len(sel), "ok": len(sel) >= need})
    return out


# ---------------------------------------------------------------- workflow
def history_before(semester_id):
    """Earlier semesters (oldest first) in the form rotation.standing() expects."""
    db = get_db()
    history = []
    for sem in db.execute("SELECT id, name FROM semesters WHERE id<? ORDER BY id", (semester_id,)).fetchall():
        rows = audit_rows(sem["id"])
        if not rows:                                  # never generated: it says nothing about who was audited
            continue
        history.append({
            "name": sem["name"],
            "present": {l["staff_id"] for l in semester_lecturers(sem["id"]) if logic.is_eligible(l["status"])},
            "audited": {a["ee_staff"] for a in rows if a["round_no"] == 1},
            "auditor_counts": dict(_count(a["er_staff"] for a in rows if a["er_staff"] and is_active(a))),
        })
    return history


def _count(items):
    counts = defaultdict(int)
    for i in items:
        counts[i] += 1
    return counts


def generate(semester, rng=None, today=None):
    """Select auditees (per-department quota, longest-waiting first) and auditors. Replaces any existing list."""
    rng = rng or random.Random()
    db = get_db()
    people = [dict(l) for l in semester_lecturers(semester["id"]) if logic.is_eligible(l["status"])]
    if not people:
        raise ValueError("No active lecturers in this semester's list.")
    by = defaultdict(list)
    for p in people:
        by[p["bidang"]].append(p)
    st = rotation.standing(people, history_before(semester["id"]), today)
    auditees, _short = rotation.select_auditees(by, semester["min_percent"], st, rng)
    pairing = rotation.assign_auditors(auditees, people, st, rng, today)
    db.execute("DELETE FROM audits WHERE semester_id=?", (semester["id"],))
    for a in auditees:
        db.execute("INSERT INTO audits (semester_id,auditee_id,auditor_id,origin) VALUES (?,?,?,'auto')",
                   (semester["id"], a["id"], pairing[a["id"]]))
    db.execute("UPDATE semesters SET state='generated', head_comment='', verified_at=NULL WHERE id=?",
               (semester["id"],))
    unmatched = sum(1 for v in pairing.values() if v is None)
    link = f"/semesters/{semester['id']}"
    for adm in users_with_role("admin"):
        tell(adm, f"Audit list generated for {semester['name']}",
             f"The audit list for {semester['name']} has been generated: {len(auditees)} auditees"
             f"{f', {unmatched} without an eligible auditor' if unmatched else ''}.\n"
             f"Please review and verify it.", link)
    db.commit()
    return len(auditees), unmatched


def submit_to_head(semester):
    """Admin verifies the list and sends it to the Head of Faculty."""
    audits = audit_rows(semester["id"])
    if not audits:
        raise ValueError("The audit list is empty.")
    if not all(d["ok"] for d in department_quota(semester)):
        raise ValueError("Some departments are below the minimum percentage.")
    db = get_db()
    db.execute("UPDATE semesters SET state='pending_head', verified_at=?, head_comment='' WHERE id=?",
               (now(), semester["id"]))
    open_rows = sum(1 for a in audits if a["auditor_id"] is None)
    for head in users_with_role("head"):
        tell(head, f"Audit list for {semester['name']} awaits your approval",
             f"The admin has verified the audit list for {semester['name']}. "
             + (f"{open_rows} auditee(s) have no auditor and need you to pick one. " if open_rows else "")
             + "Please approve it or request changes.", f"/semesters/{semester['id']}")
    db.commit()


def request_changes(semester, comment):
    db = get_db()
    db.execute("UPDATE semesters SET state='changes_requested', head_comment=? WHERE id=?",
               (comment, semester["id"]))
    for adm in users_with_role("admin"):
        tell(adm, f"Head of Faculty requested changes to {semester['name']}",
             f"Requested changes:\n{comment}", f"/semesters/{semester['id']}")
    db.commit()


def approve(semester):
    """Head approves; auditees and auditors are e-mailed."""
    if any(a["auditor_id"] is None for a in audit_rows(semester["id"])):
        raise ValueError("Pick an auditor for every auditee before approving.")
    db = get_db()
    db.execute("UPDATE semesters SET state='approved', approved_at=?, head_comment='' WHERE id=?",
               (now(), semester["id"]))
    audits = audit_rows(semester["id"])
    by_auditor = defaultdict(list)
    for a in audits:
        by_auditor[a["auditor_id"]].append(a)
        ee = db.execute("SELECT * FROM lecturers WHERE id=?", (a["auditee_id"],)).fetchone()
        tell(ee, f"You are scheduled for a teaching audit - {semester['name']}",
             f"Dear {a['ee_name']},\n\nYou have been selected for the teaching audit in {semester['name']}.\n"
             f"Your auditor is {a['er_name']} ({a['er_staff']}).\n\n"
             f"Please agree on a date and time with your auditor, then enter the session details "
             f"in the app so that reminders can be sent.", "/my-audits")
    for auditor_id, items in by_auditor.items():
        er = db.execute("SELECT * FROM lecturers WHERE id=?", (auditor_id,)).fetchone()
        names = "\n".join(f"  - {i['ee_name']} ({i['ee_staff']}, {i['ee_bidang']})" for i in items)
        tell(er, f"You are assigned as an auditor - {semester['name']}",
             f"Dear {er['name']},\n\nYou have been assigned as auditor for {semester['name']}:\n{names}\n\n"
             f"Please agree on a date and time with each auditee. Enter marks and feedback in the "
             f"app during the session.", "/my-audits")
    for adm in users_with_role("admin"):
        tell(adm, f"{semester['name']} approved; invitations sent",
             f"The Head of Faculty approved the list. {len(audits)} auditees and "
             f"{len(by_auditor)} auditors have been notified.", f"/semesters/{semester['id']}")
    db.commit()


def schedule_session(audit, session_date, session_time, venue):
    db = get_db()
    db.execute("UPDATE audits SET session_date=?, session_time=?, venue=?, reminder_sent_at=NULL WHERE id=?",
               (session_date, session_time, venue, audit["id"]))
    if audit["auditor_id"]:
        er = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditor_id"],)).fetchone()
        ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
        tell(er, f"Audit session set: {ee['name']} on {session_date}",
             f"{ee['name']} has set the audit session for {session_date} {session_time} "
             f"{('at ' + venue) if venue else ''}.", "/my-audits")
    db.commit()


HEADER_FIELDS = ("faculty", "program", "course_code", "group_no", "method")
COMMENT_FIELDS = ("comment_teaching", "comment_prof", "comment_other")


def get_scores(audit_id):
    db = get_db()
    return ({r["item"]: r["score"] for r in db.execute("SELECT item,score FROM audit_scores WHERE audit_id=?", (audit_id,))},
            {r["item"]: r["score"] for r in db.execute("SELECT item,score FROM facility_scores WHERE audit_id=?", (audit_id,))})


def build_feedback(comments):
    return "\n\n".join(f"{title}:\n{comments[k]}" for k, title in
                        [("comment_teaching", "Pengajaran dan Pembelajaran"),
                         ("comment_prof", "Profesionalisme Pensyarah"),
                         ("comment_other", "Lain-lain Perkara")] if comments[k])


def save_rubric(audit, form, submit):
    """Store the auditor's PRO-PENS rubric. With submit=True it is validated and finalised.

    Returns a list of problems (empty on success). Nothing is written if submit fails.
    """
    scores, facility = {}, {}
    for n in rubric.ITEM_NUMBERS:
        v = form.get(f"item_{n}", "")
        if v in ("1", "2", "3", "4", "5"):
            scores[n] = int(v)
    for n in range(1, len(rubric.FACILITY_ITEMS) + 1):
        v = form.get(f"fac_{n}", "")
        if v in ("0", "1", "2", "3", "4", "5"):
            facility[n] = int(v)
    header = {k: form.get(k, "").strip() for k in HEADER_FIELDS}
    comments = {k: form.get(k, "").strip() for k in COMMENT_FIELDS}
    raw_students = form.get("student_count", "").strip()
    problems = []
    students = None
    if raw_students:
        if raw_students.isdigit():
            students = int(raw_students)
        else:
            problems.append("Number of students must be a whole number.")
    if submit:
        missing = [n for n in rubric.ITEM_NUMBERS if n not in scores]
        if missing:
            problems.append("Score every item (missing: " + ", ".join(map(str, missing)) + ").")
        if not comments["comment_teaching"] or not comments["comment_prof"]:
            problems.append("Overall comments on Teaching & Learning and on Lecturer Professionalism are mandatory.")
    if problems:
        return problems

    db = get_db()
    result = rubric.compute(scores)
    db.execute("DELETE FROM audit_scores WHERE audit_id=?", (audit["id"],))
    db.executemany("INSERT INTO audit_scores VALUES (?,?,?)", [(audit["id"], n, v) for n, v in scores.items()])
    db.execute("DELETE FROM facility_scores WHERE audit_id=?", (audit["id"],))
    db.executemany("INSERT INTO facility_scores VALUES (?,?,?)", [(audit["id"], n, v) for n, v in facility.items()])
    sdate = logic.parse_date(form.get("session_date")) if form.get("session_date") else None
    db.execute(
        """UPDATE audits SET faculty=?, program=?, course_code=?, group_no=?, method=?, student_count=?,
           comment_teaching=?, comment_prof=?, comment_other=?, facility_room=?, facility_suggestion=?,
           session_date=COALESCE(?, session_date), session_time=COALESCE(NULLIF(?, ''), session_time),
           raw_score=?, marks=? WHERE id=?""",
        (header["faculty"], header["program"], header["course_code"], header["group_no"], header["method"],
         students, comments["comment_teaching"], comments["comment_prof"], comments["comment_other"],
         form.get("facility_room", "").strip(), form.get("facility_suggestion", "").strip(),
         sdate.isoformat() if sdate else None, form.get("session_time", ""),
         result["raw"] if scores else None, result["tas"] if result["complete"] else None, audit["id"]))
    if submit:
        feedback = build_feedback(comments)
        db.execute("UPDATE audits SET feedback=?, completed_at=? WHERE id=?", (feedback, now(), audit["id"]))
        ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
        tell(ee, "Feedback from your teaching audit",
             f"Dear {ee['name']},\n\nFeedback from your auditor:\n\n{feedback}\n\n"
             "Note: marks are not disclosed by the auditor. They are announced officially by the "
             "Ketua Pusat Pengajian after approval by the committee.", "/my-audits")
        sem = db.execute("SELECT * FROM semesters WHERE id=?", (audit["semester_id"],)).fetchone()
        _handle_low_mark(audit, sem, result["tas"])
        _handle_cemerlang(audit, sem, result["tas"])
    db.commit()
    return []


def is_active(audit):
    """Round-1 audits, and re-audits once the Head has approved the second auditor."""
    return audit["round_no"] == 1 or audit["reaudit_state"] == "approved"


def final_audits(audits):
    """One audit per auditee: the latest round. A second audit replaces the first mark."""
    best = {}
    for a in audits:
        if a["auditee_id"] not in best or a["round_no"] > best[a["auditee_id"]]["round_no"]:
            best[a["auditee_id"]] = a
    return list(best.values())


def _admins_and_heads():
    return users_with_role("admin"), users_with_role("head")


def _handle_low_mark(audit, semester, tas):
    """Marks below the semester threshold: alert admin + Head and propose a re-audit.

    The first audit starts a re-audit proposal (different auditor, same rule as normal matching).
    A low mark in the re-audit itself is only flagged; no third audit is created.
    """
    if tas >= semester["reaudit_below"]:
        return
    db = get_db()
    ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
    er = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditor_id"],)).fetchone()
    admins, heads = _admins_and_heads()
    link = f"/semesters/{semester['id']}"
    head_line = f"{ee['name']} ({ee['staff_id']}) scored {tas:g}/100 (below {semester['reaudit_below']:g}) " \
                f"in {semester['name']}, audited by {er['name']}."
    if audit["round_no"] >= 2:
        for p in admins + heads:
            tell(p, f"Still below {semester['reaudit_below']:g} after re-audit: {ee['name']}",
                 head_line + "\nThis was the re-audit. No further audit was created; a decision is needed.", link)
        return
    pool = [dict(l) for l in semester_lecturers(semester["id"]) if logic.is_eligible(l["status"])]
    prior = {r["auditor_id"] for r in db.execute(
        "SELECT auditor_id FROM audits WHERE semester_id=? AND auditee_id=?", (semester["id"], ee["id"]))}
    load = defaultdict(int)
    for a in audit_rows(semester["id"]):
        if is_active(a) and a["auditor_id"]:
            load[a["auditor_id"]] += 1
    pick = logic.pick_auditor(dict(ee), pool, prior | {ee["id"]}, load, random.SystemRandom())
    db.execute("""INSERT INTO audits (semester_id,auditee_id,auditor_id,origin,round_no,reaudit_of,reaudit_state)
                  VALUES (?,?,?,'auto',2,?,'proposed')""", (semester["id"], ee["id"], pick, audit["id"]))
    for p in admins:
        tell(p, f"Re-audit needed: {ee['name']} scored below {semester['reaudit_below']:g}",
             head_line + "\nA second audit by a different auditor is required this semester. "
             + ("A replacement auditor has been proposed; review it and submit it to the Head of Faculty."
                if pick else "No eligible auditor could be proposed; choose one and submit it to the Head."), link)
    for p in heads:
        tell(p, f"Low audit result: {ee['name']}",
             head_line + "\nA second audit is required. The admin will submit a proposed second auditor "
             "for your approval.", link)


def _handle_cemerlang(audit, semester, tas):
    """Final mark in the Cemerlang band: alert admin and Head (the lecturer is not told the mark)."""
    if tas < rubric.CEMERLANG_MIN:
        return
    db = get_db()
    ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
    er = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditor_id"],)).fetchone()
    admins, heads = _admins_and_heads()
    for p in admins + heads:
        tell(p, f"Cemerlang result: {ee['name']} ({tas:g}/100)",
             f"{ee['name']} ({ee['staff_id']}, {ee['bidang']}) scored {tas:g}/100 in {semester['name']} "
             f"(Cemerlang, {rubric.CEMERLANG_MIN}-100), audited by {er['name']}.", f"/semesters/{semester['id']}")


def corrections_for(audit_id):
    return get_db().execute(
        """SELECT c.*, l.name by_name FROM audit_corrections c JOIN lecturers l ON l.id=c.corrected_by
           WHERE c.audit_id=? ORDER BY c.id DESC""", (audit_id,)).fetchall()


def correct_audit(audit, admin, form):
    """Admin corrects a submitted audit's item scores and/or comments. A written reason is mandatory.

    Everything downstream follows: marks, bands, Cemerlang list, reports, Excel. Logged in audit_corrections.
    Returns a list of problems (empty on success).
    """
    db = get_db()
    if not audit["completed_at"]:
        return ["Only submitted audits can be corrected."]
    old_scores, _ = get_scores(audit["id"])
    new_scores = {}
    for n in rubric.ITEM_NUMBERS:
        v = form.get(f"item_{n}", "")
        if v in ("1", "2", "3", "4", "5"):
            new_scores[n] = int(v)
    comments = {k: form.get(k, "").strip() for k in COMMENT_FIELDS}
    reason = form.get("reason", "").strip()
    problems = []
    if len(new_scores) != len(rubric.ITEM_NUMBERS):
        problems.append("Every item must have a score from 1 to 5.")
    if not comments["comment_teaching"] or not comments["comment_prof"]:
        problems.append("Overall comments on Teaching & Learning and on Lecturer Professionalism are mandatory.")
    if not reason:
        problems.append("A reason for the correction is required.")
    changes = [f"Item {n}: {old_scores.get(n)} -> {new_scores[n]}"
               for n in rubric.ITEM_NUMBERS if n in new_scores and old_scores.get(n) != new_scores[n]]
    old_comments = {k: audit[k] for k in COMMENT_FIELDS}
    comment_changed = [k for k in COMMENT_FIELDS if old_comments[k] != comments[k]]
    changes += [f"Comment changed: {k.replace('comment_', '')}" for k in comment_changed]
    if not problems and not changes:
        problems.append("Nothing was changed.")
    if problems:
        return problems

    sem = db.execute("SELECT * FROM semesters WHERE id=?", (audit["semester_id"],)).fetchone()
    old_marks = audit["marks"]
    result = rubric.compute(new_scores)
    new_marks = result["tas"]
    db.execute("DELETE FROM audit_scores WHERE audit_id=?", (audit["id"],))
    db.executemany("INSERT INTO audit_scores VALUES (?,?,?)", [(audit["id"], n, v) for n, v in new_scores.items()])
    db.execute("""UPDATE audits SET comment_teaching=?, comment_prof=?, comment_other=?, feedback=?,
                  raw_score=?, marks=? WHERE id=?""",
               (comments["comment_teaching"], comments["comment_prof"], comments["comment_other"],
                build_feedback(comments), result["raw"], new_marks, audit["id"]))
    db.execute("INSERT INTO audit_corrections (audit_id,corrected_by,reason,old_marks,new_marks,changes,created_at) "
               "VALUES (?,?,?,?,?,?,?)",
               (audit["id"], admin["id"], reason, old_marks, new_marks, "; ".join(changes), now()))

    ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
    er = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditor_id"],)).fetchone()
    link = f"/semesters/{sem['id']}"
    fresh = db.execute("SELECT * FROM audits WHERE id=?", (audit["id"],)).fetchone()
    below, was_below = new_marks < sem["reaudit_below"], old_marks is not None and old_marks < sem["reaudit_below"]
    notes = []
    if below and not was_below:
        has_reaudit = db.execute("SELECT 1 FROM audits WHERE reaudit_of=?", (audit["id"],)).fetchone()
        if audit["round_no"] == 1 and has_reaudit:
            pass
        elif audit["round_no"] == 1 and sem["state"] != "approved":
            notes.append(f"The corrected mark is now below {sem['reaudit_below']:g} but the semester is closed, "
                         "so no re-audit was created. Decide manually.")
        else:
            _handle_low_mark(fresh, sem, new_marks)
    elif was_below and not below:
        if audit["round_no"] == 1 and db.execute("SELECT 1 FROM audits WHERE reaudit_of=?", (audit["id"],)).fetchone():
            notes.append(f"The corrected mark is no longer below {sem['reaudit_below']:g}, but a re-audit already "
                         "exists. Review it and decide whether it is still needed (it was not cancelled).")
    if new_marks >= rubric.CEMERLANG_MIN and (old_marks is None or old_marks < rubric.CEMERLANG_MIN):
        _handle_cemerlang(fresh, sem, new_marks)
    summary = (f"{ee['name']} ({ee['staff_id']}), {sem['name']}: {old_marks:g} -> {new_marks:g}. "
               f"Reason: {reason}\nChanges: {'; '.join(changes)}" + ("".join("\n" + n for n in notes)))
    admins, heads = _admins_and_heads()
    for p in admins + heads:
        tell(p, f"Audit corrected: {ee['name']}", summary, link)
    notify(er["id"], f"The admin corrected your audit of {ee['name']}", "/my-audits")
    final = audit["id"] in final_audit_ids(sem["id"])
    if comment_changed:
        tell(ee, "Your audit feedback has been updated",
             f"Dear {ee['name']},\n\nThe feedback on your audit was corrected:\n\n{build_feedback(comments)}", "/my-audits")
    if sem["released_at"] and final and old_marks != new_marks:
        tell(ee, f"Your PRO-PENS result for {sem['name']} has been updated",
             f"Dear {ee['name']},\n\nYour PRO-PENS result for {sem['name']} was corrected. "
             "Log in to My audits to view it.", "/my-audits")
        if new_marks >= rubric.CEMERLANG_MIN and (old_marks or 0) < rubric.CEMERLANG_MIN:
            notify(ee["id"], f"Congratulations! Your corrected PRO-PENS result for {sem['name']} is Cemerlang "
                             f"({rubric.CEMERLANG_MIN} and above).", "/my-audits")
    db.commit()
    return []


def final_audit_ids(semester_id):
    return {a["id"] for a in final_audits(audit_rows(semester_id))}


def release_summary(semester):
    """How many lecturers' final audits are done, for the release screen."""
    finals = final_audits(audit_rows(semester["id"]))
    return {"total": len(finals), "done": sum(1 for a in finals if a["completed_at"])}


def submit_to_ajk(semester):
    """Admin records that the results were passed to AJK Induk. Required before release."""
    if semester["state"] not in ("approved", "closed"):
        raise ValueError("The audit list must be approved first.")
    if semester["submitted_ajk_at"]:
        raise ValueError("Already recorded as submitted to AJK Induk.")
    db = get_db()
    db.execute("UPDATE semesters SET submitted_ajk_at=? WHERE id=?", (now(), semester["id"]))
    db.commit()


def release_marks(semester, allow_incomplete=False):
    """Admin releases results. Lecturers can then see their final mark; Cemerlang lecturers get an in-app alert.

    Lecturers whose final audit is not finished (e.g. re-audit pending) are not shown anything.
    """
    if not semester["submitted_ajk_at"]:
        raise ValueError("Record the submission to AJK Induk first.")
    if semester["released_at"]:
        raise ValueError("Marks are already released.")
    summary = release_summary(semester)
    if summary["done"] < summary["total"] and not allow_incomplete:
        raise ValueError(f"{summary['total'] - summary['done']} audit(s) are not finished. Tick the box to release "
                         "anyway; those lecturers will see no result.")
    db = get_db()
    db.execute("UPDATE semesters SET released_at=? WHERE id=?", (now(), semester["id"]))
    for a in final_audits(audit_rows(semester["id"])):
        if not a["completed_at"] or a["marks"] is None:
            continue
        ee = db.execute("SELECT * FROM lecturers WHERE id=?", (a["auditee_id"],)).fetchone()
        subject = f"Your PRO-PENS result for {semester['name']} is available"
        body = f"Dear {ee['name']},\n\nYour PRO-PENS result for {semester['name']} has been released. Log in to My audits to view it."
        if a["marks"] >= rubric.CEMERLANG_MIN:
            notify(ee["id"], f"Congratulations! Your PRO-PENS result for {semester['name']} is Cemerlang "
                             f"({rubric.CEMERLANG_MIN} and above).", "/my-audits")
            send_email(ee["email"], subject, body)     # in-app message above is the special one
        else:
            tell(ee, subject, body, "/my-audits")
    for head in users_with_role("head"):
        notify(head["id"], f"Marks for {semester['name']} have been released to lecturers", f"/semesters/{semester['id']}")
    db.commit()


def reaudit_rows(semester_id):
    return [a for a in audit_rows(semester_id) if a["round_no"] == 2]


def first_audit(audit):
    return get_db().execute("SELECT * FROM audits WHERE id=?", (audit["reaudit_of"],)).fetchone()


def submit_reaudit(audit, auditor_id):
    """Admin sets the second auditor and sends it to the Head for approval."""
    db = get_db()
    if audit["round_no"] != 2 or audit["reaudit_state"] not in ("proposed", "changes_requested"):
        raise ValueError("This re-audit is not open for editing.")
    prior = {r["auditor_id"] for r in db.execute(
        "SELECT auditor_id FROM audits WHERE semester_id=? AND auditee_id=? AND round_no=1",
        (audit["semester_id"], audit["auditee_id"]))}
    ok = db.execute(
        """SELECT l.* FROM semester_lecturers sl JOIN lecturers l ON l.id=sl.lecturer_id
           WHERE sl.semester_id=? AND l.id=? AND lower(l.status)='active'""",
        (audit["semester_id"], auditor_id)).fetchone()
    if not ok or auditor_id == audit["auditee_id"]:
        raise ValueError("Choose an active lecturer other than the auditee.")
    if auditor_id in prior:
        raise ValueError("The re-audit must be done by a different auditor from the first audit.")
    sem = db.execute("SELECT * FROM semesters WHERE id=?", (audit["semester_id"],)).fetchone()
    ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
    db.execute("UPDATE audits SET auditor_id=?, origin='manual', reaudit_state='pending_head', "
               "reaudit_comment='' WHERE id=?", (auditor_id, audit["id"]))
    for head in users_with_role("head"):
        tell(head, f"Approve second auditor for {ee['name']} ({sem['name']})",
             f"Re-audit for {ee['name']}: proposed auditor is {ok['name']} ({ok['grade']}). "
             "Please approve or ask the admin to change it.", f"/semesters/{sem['id']}")
    db.commit()


def decide_reaudit(audit, approve_it, comment=""):
    """Head approves the second auditor (emails go out) or sends it back."""
    db = get_db()
    if audit["reaudit_state"] != "pending_head" or not audit["auditor_id"]:
        raise ValueError("This re-audit is not awaiting approval.")
    sem = db.execute("SELECT * FROM semesters WHERE id=?", (audit["semester_id"],)).fetchone()
    ee = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditee_id"],)).fetchone()
    er = db.execute("SELECT * FROM lecturers WHERE id=?", (audit["auditor_id"],)).fetchone()
    link = f"/semesters/{sem['id']}"
    if approve_it:
        db.execute("UPDATE audits SET reaudit_state='approved', reaudit_approved_at=? WHERE id=?",
                   (now(), audit["id"]))
        tell(ee, f"A second teaching audit is required - {sem['name']}",
             f"Dear {ee['name']},\n\nA second audit is required for you in {sem['name']}. "
             f"Your auditor is {er['name']} ({er['staff_id']}). Please agree on a date and time and "
             "enter the session details in the app.", "/my-audits")
        tell(er, f"You are assigned a re-audit - {sem['name']}",
             f"Dear {er['name']},\n\nYou have been assigned to audit {ee['name']} ({ee['staff_id']}) "
             f"in {sem['name']}. Please agree on a date and time with the auditee. Enter the rubric "
             "in the app during the session.", "/my-audits")
        for adm in users_with_role("admin"):
            tell(adm, f"Re-audit approved for {ee['name']}", f"Auditor: {er['name']}. Both parties were e-mailed.", link)
    else:
        db.execute("UPDATE audits SET reaudit_state='changes_requested', reaudit_comment=? WHERE id=?",
                   (comment, audit["id"]))
        for adm in users_with_role("admin"):
            tell(adm, f"Head asked to change the second auditor for {ee['name']}", f"Reason: {comment}", link)
    db.commit()


def send_reminders(today=None):
    """Remind both parties on the day before and the day of the session (once per schedule)."""
    today = today or date.today()
    db = get_db()
    due = db.execute(
        """SELECT * FROM audits WHERE session_date IS NOT NULL AND completed_at IS NULL
           AND reminder_sent_at IS NULL AND (round_no=1 OR reaudit_state='approved')
           AND session_date BETWEEN ? AND ?""",
        (today.isoformat(), (today + timedelta(days=1)).isoformat())).fetchall()
    for a in due:
        when = f"{a['session_date']} {a['session_time'] or ''} {('at ' + a['venue']) if a['venue'] else ''}"
        for lid in (a["auditee_id"], a["auditor_id"]):
            if lid:
                p = db.execute("SELECT * FROM lecturers WHERE id=?", (lid,)).fetchone()
                tell(p, "Reminder: teaching audit session", f"Reminder: audit session on {when}.", "/my-audits")
        db.execute("UPDATE audits SET reminder_sent_at=? WHERE id=?", (now(), a["id"]))
    db.commit()
    return len(due)


# ---------------------------------------------------------------- reports
def _band(marks):
    return next(label for label, floor in BANDS if marks >= floor)


def _avg(vals):
    return round(sum(vals) / len(vals), 2) if vals else None


def report(semester):
    """Semester report. A lecturer's final mark is that of their latest audit (a re-audit replaces the first)."""
    allrows = audit_rows(semester["id"])
    audits = final_audits(allrows)                       # one row per auditee
    done = [a for a in audits if a["completed_at"]]
    scored = [a["marks"] for a in done if a["marks"] is not None]
    by_campus, by_dept = defaultdict(list), defaultdict(list)
    for a in done:
        by_campus[a["ee_campus"] or "-"].append(a["marks"])
        by_dept[a["ee_bidang"]].append(a["marks"])
    load = defaultdict(lambda: [0, 0])
    for a in allrows:
        if a["auditor_id"] and is_active(a):
            key = (a["er_staff"], a["er_name"])
            load[key][0] += 1
            load[key][1] += 1 if a["completed_at"] else 0
    bands = {label: 0 for label, _ in BANDS}
    for m in scored:
        bands[_band(m)] += 1
    sec_totals = defaultdict(list)
    for a in done:
        sc, _ = get_scores(a["id"])
        for k, v in rubric.compute(sc)["sections"].items():
            sec_totals[k].append(v)
    sections = [{"key": sec["key"], "title": sec["title"].split(":", 1)[1].strip(), "max": sec["max"],
                 "average": _avg(sec_totals.get(sec["key"], [])),
                 "percent": round(100 * _avg(sec_totals[sec["key"]]) / sec["max"], 1)
                 if sec_totals.get(sec["key"]) else None} for sec in rubric.SECTIONS]
    quota = department_quota(semester)
    for q in quota:
        q["completed"] = sum(1 for a in done if a["ee_bidang"] == q["dept"])
        q["average"] = _avg(by_dept.get(q["dept"], []))
    re = [a for a in allrows if a["round_no"] == 2]
    statuses = rule_statuses(semester["id"])
    return {
        "total_lecturers": len(semester_lecturers(semester["id"])),
        "eligible": sum(q["total"] for q in quota),
        "audits": len(audits), "completed": len(done),
        "scheduled": sum(1 for a in audits if a["session_date"] and not a["completed_at"]),
        "not_scheduled": sum(1 for a in audits if not a["session_date"] and not a["completed_at"]),
        "average": _avg(scored),
        "highest": max(scored) if scored else None,
        "lowest": min(scored) if scored else None,
        "threshold": semester["reaudit_below"],
        "reaudits": len(re), "reaudits_done": sum(1 for a in re if a["completed_at"]),
        "still_below": sum(1 for a in done if a["marks"] is not None and a["marks"] < semester["reaudit_below"]),
        "cemerlang": [{"staff_id": a["ee_staff"], "name": a["ee_name"], "bidang": a["ee_bidang"],
                       "campus": a["ee_campus"], "marks": a["marks"]}
                      for a in sorted(done, key=lambda x: -(x["marks"] or 0))
                      if a["marks"] is not None and a["marks"] >= rubric.CEMERLANG_MIN],
        "quota": quota,
        "campus": [{"campus": c, "count": len(v), "average": _avg(v)} for c, v in sorted(by_campus.items())],
        "bands": bands, "sections": sections,
        "workload": [{"staff_id": k[0], "name": k[1], "assigned": v[0], "completed": v[1]}
                     for k, v in sorted(load.items(), key=lambda kv: -kv[1][0])],
        "rule_violations": sum(1 for v in statuses.values() if v == "override"),
        "peer_pairs": sum(1 for v in statuses.values() if v == "peer"),
    }


def export_workbook(semester):
    """Marks for all audited lecturers + summary sheets. Feedback is intentionally excluded."""
    audits = audit_rows(semester["id"])
    rep = report(semester)
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Marks"
    keys = [sec["key"] for sec in rubric.SECTIONS]
    ws.append(["Semester", "Staff ID", "Staff Name", "Gred Jawatan", "Bidang", "Campus",
               "Auditor Staff ID", "Auditor Name", "Session Date", "Session Time", "Venue",
               "Course Code", "Programme", "Group", "No. Students", "Round", "Counts as final mark"]
              + [f"{k} (/{sec['max']})" for k, sec in zip(keys, rubric.SECTIONS)]
              + ["Raw Total (/145)", "Total Score TS (/50)", "Marks TAS (/100)", "Indicator", "Status"])
    final_ids = {a["id"] for a in final_audits(audits)}
    for a in audits:
        sc, _ = get_scores(a["id"])
        r = rubric.compute(sc)
        done = bool(a["completed_at"])
        ws.append([semester["name"], a["ee_staff"], a["ee_name"], a["ee_grade"], a["ee_bidang"],
                   a["ee_campus"], a["er_staff"], a["er_name"], a["session_date"], a["session_time"],
                   a["venue"], a["course_code"], a["program"], a["group_no"], a["student_count"],
                   a["round_no"], "Yes" if a["id"] in final_ids else "No (replaced by re-audit)"]
                  + [r["sections"][k] if done else None for k in keys]
                  + [r["raw"] if done else None, r["ts"] if done else None, a["marks"] if done else None,
                     rubric.indicator(a["marks"]) if done else "",
                     "Completed" if done else (
                         "Re-audit awaiting approval" if a["round_no"] == 2 and a["reaudit_state"] != "approved"
                         else "Scheduled" if a["session_date"] else "Not scheduled")])
    f = wb.create_sheet("Facilities")
    f.append(["Staff ID", "Staff Name", "Room"] + [f"{n}" for n in range(1, len(rubric.FACILITY_ITEMS) + 1)]
             + ["Suggestions"])
    for a in audits:
        if a["completed_at"]:
            _, fs = get_scores(a["id"])
            f.append([a["ee_staff"], a["ee_name"], a["facility_room"]]
                     + [fs.get(n) for n in range(1, len(rubric.FACILITY_ITEMS) + 1)] + [a["facility_suggestion"]])
    s = wb.create_sheet("Summary by Department")
    s.append(["Bidang", "Eligible lecturers", f"Required (min {semester['min_percent']:g}%)",
              "Selected", "Completed", "Average marks"])
    for q in rep["quota"]:
        s.append([q["dept"], q["total"], q["required"], q["selected"], q["completed"], q["average"]])
    c = wb.create_sheet("Summary by Campus")
    c.append(["Campus", "Audits completed", "Average marks"])
    for r in rep["campus"]:
        c.append([r["campus"], r["count"], r["average"]])
    w = wb.create_sheet("Auditor Workload")
    w.append(["Staff ID", "Auditor", "Assigned", "Completed"])
    for r in rep["workload"]:
        w.append([r["staff_id"], r["name"], r["assigned"], r["completed"]])
    cm = wb.create_sheet("Cemerlang")
    cm.append(["Staff ID", "Staff Name", "Bidang", "Campus", "Marks TAS (/100)", "Indicator"])
    for r in rep["cemerlang"]:
        cm.append([r["staff_id"], r["name"], r["bidang"], r["campus"], r["marks"], rubric.indicator(r["marks"])])
    co = wb.create_sheet("Corrections")
    co.append(["Staff ID", "Staff Name", "Round", "Corrected on", "Corrected by", "Old marks", "New marks",
               "Reason", "Changes"])
    for a in audits:
        for c in reversed(corrections_for(a["id"])):
            co.append([a["ee_staff"], a["ee_name"], a["round_no"], c["created_at"], c["by_name"],
                       c["old_marks"], c["new_marks"], c["reason"], c["changes"]])
    sa = wb.create_sheet("Section Averages")
    sa.append(["Section", "Max", "Average", "Average %"])
    for r in rep["sections"]:
        sa.append([f"{r['key']}: {r['title']}", r["max"], r["average"], r["percent"]])
    b = wb.create_sheet("Mark Bands")
    b.append(["Band", "Lecturers"])
    for k, v in rep["bands"].items():
        b.append([k, v])
    fill = PatternFill("solid", fgColor="DDE6F3")
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 40)
        sheet.freeze_panes = "A2"
    return wb
