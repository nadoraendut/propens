import io
import sqlite3

import pytest
from openpyxl import load_workbook

from propens import create_app, rubric, services
from propens.db import get_db
from test_workflow import app, client_for, post  # noqa: F401  (fixture + helpers)

CSV = ("Staff ID,Staff Name,Gred Jawatan,Bidang,Start of Employment,Campus,Status,Email\n"
       "S1,L1,DS54,CS,1990-01-01,N,Active,s1@u.edu\nS7,L7,DS54,CS,1990-01-01,N,Active,s7@u.edu\n"
       "S2,L2,DS52,CS,1995-01-01,N,Active,s2@u.edu\nS3,L3,DS48,CS,2000-01-01,S,Active,s3@u.edu\n"
       "S4,L4,DS45,CS,2005-01-01,S,Active,s4@u.edu\nS5,L5,DS44,CS,2010-01-01,N,Active,s5@u.edu\n"
       "S6,L6,DS41,CS,2015-01-01,S,Active,s6@u.edu\n")


def score_all(c, audit_id, value, **extra):
    data = {f"item_{n}": str(value) for n in rubric.ITEM_NUMBERS}
    return post(c, f"/audits/{audit_id}/rubric", action="submit", comment_teaching="ok", comment_prof="ok",
                **data, **extra)


@pytest.fixture
def approved(app):
    admin, head = client_for(app, "ADM"), client_for(app, "HOF")
    post(admin, "/semesters/new", name="S", min_percent="20")
    with admin.session_transaction() as s:
        s["csrf"] = "t"
    admin.post("/semesters/1/upload", content_type="multipart/form-data",
               data={"csrf": "t", "file": (io.BytesIO(CSV.encode()), "l.csv")})
    post(admin, "/semesters/1/generate")
    post(admin, "/semesters/1/verify")
    post(head, "/semesters/1/decision", action="approve")
    with app.app_context():
        first = [a for a in services.audit_rows(1) if a["round_no"] == 1][0]
        assert get_db().execute("SELECT reaudit_below FROM semesters").fetchone()[0] == 80
    return admin, head, dict(first)


def mails(app, to, like):
    with app.app_context():
        return get_db().execute("SELECT * FROM emails WHERE to_addr=? AND subject LIKE ?", (to, like)).fetchall()


def test_exactly_80_does_not_trigger(app, approved):
    admin, head, first = approved
    score_all(client_for(app, first["er_staff"], first["er_staff"]), first["id"], 4)   # TAS 80.0
    with app.app_context():
        assert get_db().execute("SELECT marks FROM audits WHERE id=?", (first["id"],)).fetchone()[0] == 80.0
        assert not services.reaudit_rows(1)


def low_then_approve(app, approved):
    """Low first mark -> proposal -> admin submits -> Head sends back once -> admin resubmits -> Head approves."""
    admin, head, first = approved
    er1 = client_for(app, first["er_staff"], first["er_staff"])
    score_all(er1, first["id"], 2)                                         # TAS 40.0 -> below 80
    with app.app_context():
        db = get_db()
        re = services.reaudit_rows(1)
        assert len(re) == 1 and re[0]["reaudit_state"] == "proposed" and re[0]["round_no"] == 2
        assert re[0]["auditor_id"] not in (first["auditor_id"], first["auditee_id"])
        rid = re[0]["id"]
        # admin and Head alerted (generated list + verify already notified once each)
        assert db.execute("SELECT COUNT(*) FROM notifications WHERE lecturer_id=1 AND message LIKE 'Re-audit needed%'").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM notifications WHERE lecturer_id=2 AND message LIKE 'Low audit result%'").fetchone()[0] == 1
        options = [l for l in services.semester_lecturers(1)
                   if l["id"] not in (first["auditor_id"], first["auditee_id"])]
    # nothing reaches the auditee until the Head approves; the re-audit is not yet usable
    assert not mails(app, first["ee_email"], "A second teaching audit%")
    post(client_for(app, first["ee_staff"], first["ee_staff"]), f"/audits/{rid}/schedule", session_date="2026-11-01")
    with app.app_context():
        assert get_db().execute("SELECT session_date FROM audits WHERE id=?", (rid,)).fetchone()[0] is None

    # different-auditor rule
    r = post(admin, f"/audits/{rid}/reaudit/submit", auditor_id=first["auditor_id"])
    assert b"different auditor" in r.data
    new_auditor = options[0]
    post(admin, f"/audits/{rid}/reaudit/submit", auditor_id=new_auditor["id"])
    # head must approve: send back needs a reason, then admin resubmits
    assert b"say what should change" in post(head, f"/audits/{rid}/reaudit/decision", action="changes").data
    post(head, f"/audits/{rid}/reaudit/decision", action="changes", comment="Prefer someone else")
    with app.app_context():
        assert get_db().execute("SELECT reaudit_state FROM audits WHERE id=?", (rid,)).fetchone()[0] == "changes_requested"
    assert not mails(app, first["ee_email"], "A second teaching audit%")
    post(admin, f"/audits/{rid}/reaudit/submit", auditor_id=options[1]["id"])
    post(head, f"/audits/{rid}/reaudit/decision", action="approve")
    chosen = options[1]
    assert mails(app, first["ee_email"], "A second teaching audit%")
    assert mails(app, chosen["email"], "You are assigned a re-audit%")
    return admin, head, first, rid, chosen


def test_reaudit_replaces_first_mark(app, approved):
    admin, head, first, rid, chosen = low_then_approve(app, approved)
    assert b"re-audits before closing" in post(admin, "/semesters/1/close").data
    # the new auditor (not the first one) can now score; the first auditor cannot
    assert client_for(app, first["er_staff"], first["er_staff"]).get(f"/audits/{rid}/rubric").status_code == 403
    er2 = client_for(app, chosen["staff_id"], chosen["staff_id"])
    score_all(er2, rid, 5)                                                 # TAS 100
    with app.app_context():
        rep = services.report(get_db().execute("SELECT * FROM semesters").fetchone())
        # only this lecturer is scored: their final mark is the re-audit's 100, not the first 40, not an average
        assert rep["completed"] == 1 and rep["average"] == 100.0 and rep["lowest"] == 100.0
        assert rep["reaudits"] == 1 and rep["reaudits_done"] == 1 and rep["still_below"] == 0
        assert len(services.reaudit_rows(1)) == 1                          # no third audit
    ws = load_workbook(io.BytesIO(admin.get("/semesters/1/export.xlsx").data))["Audit Marks"]
    hdr = [c.value for c in ws[1]]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index("Staff ID")] == first["ee_staff"]]
    assert sorted(r[hdr.index("Round")] for r in rows) == [1, 2]
    final = {r[hdr.index("Round")]: r[hdr.index("Counts as final mark")] for r in rows}
    assert final[2] == "Yes" and final[1].startswith("No")
    assert post(admin, "/semesters/1/close").status_code == 200
    with app.app_context():
        assert get_db().execute("SELECT state FROM semesters").fetchone()[0] == "closed"


def test_second_low_mark_is_flagged_not_repeated(app, approved):
    admin, head, first, rid, chosen = low_then_approve(app, approved)
    score_all(client_for(app, chosen["staff_id"], chosen["staff_id"]), rid, 2)   # still 40.0
    with app.app_context():
        db = get_db()
        assert len(services.reaudit_rows(1)) == 1
        assert max(r["round_no"] for r in services.audit_rows(1)) == 2
        for uid in (1, 2):
            assert db.execute("SELECT COUNT(*) FROM notifications WHERE lecturer_id=? AND message LIKE 'Still below%'",
                              (uid,)).fetchone()[0] == 1
        rep = services.report(db.execute("SELECT * FROM semesters").fetchone())
        assert rep["still_below"] == 1
    # feedback to the auditee still contains no marks
    fb = mails(app, first["ee_email"], "Feedback%")
    assert len(fb) == 2 and all("40" not in m["body"] for m in fb)


def test_old_database_is_upgraded_and_keeps_data(tmp_path):
    path = str(tmp_path / "old.db")
    con = sqlite3.connect(path)
    con.executescript("""
    CREATE TABLE lecturers (id INTEGER PRIMARY KEY AUTOINCREMENT, staff_id TEXT NOT NULL UNIQUE, name TEXT NOT NULL,
      grade TEXT NOT NULL, bidang TEXT NOT NULL, start_date TEXT NOT NULL, campus TEXT NOT NULL DEFAULT '',
      status TEXT NOT NULL DEFAULT 'Active', email TEXT NOT NULL DEFAULT '', role TEXT NOT NULL DEFAULT 'lecturer',
      password_hash TEXT NOT NULL);
    CREATE TABLE semesters (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL UNIQUE, min_percent REAL NOT NULL,
      state TEXT NOT NULL DEFAULT 'draft', verified_at TEXT, approved_at TEXT, head_comment TEXT NOT NULL DEFAULT '',
      created_at TEXT NOT NULL);
    CREATE TABLE audits (id INTEGER PRIMARY KEY AUTOINCREMENT, semester_id INTEGER NOT NULL REFERENCES semesters(id),
      auditee_id INTEGER NOT NULL REFERENCES lecturers(id), auditor_id INTEGER REFERENCES lecturers(id),
      origin TEXT NOT NULL DEFAULT 'auto', session_date TEXT, session_time TEXT, venue TEXT NOT NULL DEFAULT '',
      reminder_sent_at TEXT, marks REAL, feedback TEXT NOT NULL DEFAULT '', completed_at TEXT,
      UNIQUE (semester_id, auditee_id));
    INSERT INTO lecturers (staff_id,name,grade,bidang,start_date,password_hash) VALUES ('A','A','DS1','CS','2000-01-01','x'),('B','B','DS2','CS','2001-01-01','x');
    INSERT INTO semesters (name,min_percent,created_at) VALUES ('Old',20,'2026-01-01');
    INSERT INTO audits (semester_id,auditee_id,auditor_id,marks,completed_at) VALUES (1,1,2,77.5,'2026-02-02');
    """)
    con.commit(); con.close()
    app = create_app({"DATABASE": path, "TESTING": True})
    with app.app_context():
        db = get_db()
        row = db.execute("SELECT * FROM audits").fetchone()
        assert row["marks"] == 77.5 and row["round_no"] == 1
        assert db.execute("SELECT reaudit_below FROM semesters").fetchone()[0] == 80
        db.execute("INSERT INTO audits (semester_id,auditee_id,auditor_id,round_no) VALUES (1,1,2,2)")  # now allowed
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("INSERT INTO audits (semester_id,auditee_id,auditor_id,round_no) VALUES (1,1,2,2)")


def test_90_and_above_is_cemerlang(app, approved):
    admin, head, first = approved
    score_all(client_for(app, first["er_staff"], first["er_staff"]), first["id"], 5)          # TAS 100
    with app.app_context():
        sem = get_db().execute("SELECT * FROM semesters").fetchone()
        rep = services.report(sem)
        assert [c["staff_id"] for c in rep["cemerlang"]] == [first["ee_staff"]]
        assert rubric.indicator(90) .startswith("Cemerlang") and rubric.indicator(89.99).startswith("Baik")
    assert b"Cemerlang" in admin.get("/semesters/1").data and b"Cemerlang" in admin.get("/semesters/1/report").data
    wb = load_workbook(io.BytesIO(admin.get("/semesters/1/export.xlsx").data))
    assert wb["Cemerlang"]["A2"].value == first["ee_staff"] and wb["Cemerlang"]["F2"].value.startswith("Cemerlang")
    # the lecturer being rated never sees it
    ee = client_for(app, first["ee_staff"], first["ee_staff"]).get("/my-audits").data.decode()
    assert "Cemerlang" not in ee.split("Audits where I am the auditor")[0]
