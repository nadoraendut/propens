import io
import random
import re
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from propens import create_app, rubric, services
from propens.db import get_db


@pytest.fixture
def app(tmp_path):
    app = create_app({"DATABASE": str(tmp_path / "t.db"), "TESTING": True})
    with app.app_context():
        from werkzeug.security import generate_password_hash
        db = get_db()
        for sid, name, role, email in [("ADM", "Admin", "admin", "admin@u.edu"),
                                       ("HOF", "Head", "head", "head@u.edu")]:
            db.execute("""INSERT INTO lecturers (staff_id,name,grade,bidang,start_date,email,role,password_hash)
                          VALUES (?,?,?,?,?,?,?,?)""",
                       (sid, name, "DS54", "Admin", "1995-01-01", email, role, generate_password_hash("pw123456")))
        db.commit()
    return app


def client_for(app, staff_id, pw="pw123456"):
    c = app.test_client()
    c.get("/login")
    with c.session_transaction() as s:
        s["csrf"] = "t"
    r = c.post("/login", data={"staff_id": staff_id, "password": pw, "csrf": "t"})
    assert r.status_code == 302
    return c


def post(c, url, **data):
    with c.session_transaction() as s:
        s["csrf"] = "t"
    return c.post(url, data={**data, "csrf": "t"}, follow_redirects=True)


def make_csv():
    rows = ["Staff ID,Staff Name,Gred Jawatan,Bidang,Start of Employment,Campus,Status,Email"]
    n = 0
    for dept, size in [("CS", 10), ("Math", 5), ("Physics", 4)]:
        for i in range(size):
            n += 1
            rows.append(f"S{n:03d},Lect {n},DS{41 + (n % 4) * 3},{dept},{2000 + n}-01-01,"
                        f"{'North' if n % 2 else 'South'},Active,s{n}@u.edu")
    rows.append("S900,On Leave,DS52,CS,2001-01-01,North,Study Leave,x@u.edu")
    rows.append("S901,Bad Date,DS52,CS,notadate,North,Active,y@u.edu")
    return "\n".join(rows).encode()


def test_full_workflow(app):
    admin, head = client_for(app, "ADM"), client_for(app, "HOF")
    post(admin, "/semesters/new", name="Sem 1 2026", min_percent="20")

    # upload: bad row reported, good rows kept
    with admin.session_transaction() as s:
        s["csrf"] = "t"
    r = admin.post("/semesters/1/upload", follow_redirects=True, content_type="multipart/form-data",
                   data={"csrf": "t", "file": (io.BytesIO(make_csv()), "l.csv")})
    assert b"Imported 20 lecturer" in r.data and b"Row 22" in r.data

    # generate: quota 20% -> CS ceil(2.0)=2, Math 1, Physics 1
    post(admin, "/semesters/1/generate")
    with app.app_context():
        sem = get_db().execute("SELECT * FROM semesters WHERE id=1").fetchone()
        assert sem["state"] == "generated"
        quota = {q["dept"]: q for q in services.department_quota(sem)}
        assert (quota["CS"]["selected"], quota["Math"]["selected"], quota["Physics"]["selected"]) == (2, 1, 1)
        assert all(q["ok"] for q in quota.values())
        audits = services.audit_rows(1)
        assert all(v in ('ok', 'peer', 'none') for v in services.rule_statuses(1).values())
        assert not any(a["ee_staff"] == "S900" for a in audits)      # not Active
        assert get_db().execute("SELECT COUNT(*) FROM notifications WHERE lecturer_id=1").fetchone()[0] == 1

    # admin verifies -> head alerted; list locked
    post(admin, "/semesters/1/verify")
    r = post(admin, "/semesters/1/generate")
    assert b"locked" in r.data
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) FROM notifications WHERE lecturer_id=2").fetchone()[0] == 1

    # head requests changes, admin edits, re-verifies, head approves
    post(head, "/semesters/1/decision", action="changes", comment="Remove one Math person")
    with app.app_context():
        assert get_db().execute("SELECT state FROM semesters").fetchone()[0] == "changes_requested"
        math = [a for a in services.audit_rows(1) if a["ee_bidang"] == "Math"][0]
    r = post(admin, f"/audits/{math['id']}/delete")
    r = post(admin, "/semesters/1/verify")
    assert b"below the minimum" in r.data                               # quota guard
    post(admin, "/semesters/1/audits/add", auditee_id=math["auditee_id"], auditor_id=math["auditor_id"])
    post(admin, "/semesters/1/verify")
    post(head, "/semesters/1/decision", action="approve")

    with app.app_context():
        db = get_db()
        assert db.execute("SELECT state FROM semesters").fetchone()[0] == "approved"
        audits = services.audit_rows(1)
        mails = {m["to_addr"] for m in db.execute("SELECT * FROM emails")}
        for a in audits:
            assert a["ee_email"] in mails and a["er_email"] in mails
        a = audits[0]

    # auditee schedules; auditor is notified; reminder fires once
    ee = client_for(app, a["ee_staff"], a["ee_staff"])
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    post(ee, f"/audits/{a['id']}/schedule", session_date=tomorrow, session_time="10:00", venue="Room 1")
    with app.app_context():
        assert services.send_reminders() == 1
        assert services.send_reminders() == 0

    # only the auditor may open/submit the rubric; the auditee may not even view it
    assert ee.get(f"/audits/{a['id']}/rubric").status_code == 403
    assert ee.post(f"/audits/{a['id']}/rubric", data={"csrf": "t", "action": "submit"}).status_code == 403

    # auditor: incomplete submission rejected, nothing stored as completed
    er = client_for(app, a["er_staff"], a["er_staff"])
    r = post(er, f"/audits/{a['id']}/rubric", action="submit", item_1="5", comment_teaching="x")
    assert b"Score every item" in r.data and b"mandatory" in r.data
    # draft keeps partial scores
    post(er, f"/audits/{a['id']}/rubric", action="save", item_1="5", item_2="3", faculty="FSKM")
    page = er.get(f"/audits/{a['id']}/rubric").data.decode()
    assert 'value="FSKM"' in page
    with app.app_context():
        assert get_db().execute("SELECT completed_at, marks FROM audits WHERE id=?", (a["id"],)).fetchone()[:] == (None, None)

    # full submission: all 4s => raw 116 => TAS 80.0
    full = {f"item_{n}": "4" for n in rubric.ITEM_NUMBERS}
    full.update({"fac_1": "5", "fac_3": "0", "facility_room": "BK1", "facility_suggestion": "More plugs"})
    post(er, f"/audits/{a['id']}/rubric", action="submit", comment_teaching="Clear delivery.",
         comment_prof="Punctual.", course_code="CSC101", student_count="40", **full)
    with app.app_context():
        row = get_db().execute("SELECT * FROM audits WHERE id=?", (a["id"],)).fetchone()
        assert row["marks"] == 80.0 and row["raw_score"] == 116 and row["completed_at"]
        fb = get_db().execute("SELECT * FROM emails WHERE to_addr=? AND subject LIKE 'Feedback%'",
                              (a["ee_email"],)).fetchone()
        assert "Clear delivery." in fb["body"] and "Punctual." in fb["body"]
        assert "80" not in fb["body"].replace("Ketua", "") and "More plugs" not in fb["body"]
    # locked after submit
    assert er.post(f"/audits/{a['id']}/rubric", data={"csrf": "t", "action": "save"}).status_code == 403
    page = ee.get("/my-audits").data.decode()
    assert "Clear delivery." in page and "80.0" not in page and "rubric" not in page.lower().replace("pro-pens rubric", "")
    assert admin.get(f"/audits/{a['id']}/rubric").status_code == 200
    assert b"80.0" in admin.get(f"/audits/{a['id']}/rubric").data

    # Excel export: marks + section totals, no free-text comments
    r = admin.get("/semesters/1/export.xlsx")
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb["Audit Marks"]
    hdr = [c.value for c in ws[1]]
    row = [c.value for c in next(x for x in ws.iter_rows(min_row=2) if x[hdr.index("Marks TAS (/100)")].value is not None)]
    assert row[hdr.index("Marks TAS (/100)")] == 80.0 and row[hdr.index("Raw Total (/145)")] == 116
    assert row[hdr.index("Indicator")].startswith("Baik") and row[hdr.index("A (/25)")] == 20
    assert "Clear delivery." not in str([[c.value for c in r_] for r_ in ws.iter_rows()])
    assert {"Summary by Department", "Auditor Workload", "Facilities", "Section Averages"} <= set(wb.sheetnames)
    assert wb["Facilities"]["D2"].value == 5

    assert admin.get("/semesters/1/report").status_code == 200
    assert ee.get("/semesters/1/report").status_code == 403


def test_access_control_and_csrf(app):
    c = app.test_client()
    assert c.get("/semesters/1").status_code == 302                     # not logged in
    assert c.post("/login", data={"staff_id": "ADM", "password": "pw123456"}).status_code == 400
    head = client_for(app, "HOF")
    assert head.get("/lecturers").status_code == 403


def test_head_picks_auditor_for_most_senior(app):
    admin, head = client_for(app, "ADM"), client_for(app, "HOF")
    post(admin, "/semesters/new", name="S", min_percent="100")
    csv_ = ("Staff ID,Staff Name,Gred Jawatan,Bidang,Start of Employment,Campus,Status,Email\n"
            "A,Senior,DS54,CS,1990-01-01,N,Active,a@u.edu\nB,Junior,DS45,CS,2015-01-01,N,Active,b@u.edu\n")
    with admin.session_transaction() as s:
        s["csrf"] = "t"
    admin.post("/semesters/1/upload", content_type="multipart/form-data",
               data={"csrf": "t", "file": (io.BytesIO(csv_.encode()), "l.csv")})
    post(admin, "/semesters/1/generate")
    r = post(admin, "/semesters/1/verify")
    assert b"sent to the Head" in r.data
    with app.app_context():
        senior = [a for a in services.audit_rows(1) if a["ee_staff"] == "A"][0]
        assert senior["auditor_id"] is None
        b_id = get_db().execute("SELECT id FROM lecturers WHERE staff_id='B'").fetchone()[0]
    r = post(head, "/semesters/1/decision", action="approve")
    assert b"Pick an auditor" in r.data
    post(head, f"/audits/{senior['id']}/head-assign", auditor_id=b_id)
    post(head, "/semesters/1/decision", action="approve")
    with app.app_context():
        assert get_db().execute("SELECT state FROM semesters").fetchone()[0] == "approved"
