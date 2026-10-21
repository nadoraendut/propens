"""Read the history ledger (the file this tool produced last semester, or a simple list of past audits)."""
import csv
import io
import re

from openpyxl import load_workbook

ALIASES = {
    "semester": ["semester", "sem", "semesters", "sesi"],
    "staff_id": ["staffid", "id", "staffno", "nopekerja"],
    "name": ["staffname", "name", "nama"],
    "bidang": ["bidang", "department", "field"],
    "in_pool": ["inpool", "pool", "inthepool"],
    "auditee": ["auditee", "audited", "diaudit"],
    "auditor_count": ["auditorcount", "auditsconducted", "served", "auditor"],
}


def _norm(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def read_ledger(filename, data):
    """Return canonical ledger rows: dicts with semester, staff_id (+ optional name, bidang, in_pool,
    auditee, auditor_count). In an .xlsx the sheet called "Ledger" is used if present."""
    low = filename.lower()
    if low.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            raise ValueError("History: could not read the file as an Excel (.xlsx) workbook. Is it damaged, "
                             "or a CSV saved with an .xlsx name?") from None
        ws = wb["Ledger"] if "Ledger" in wb.sheetnames else wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    elif low.endswith(".csv"):
        raw = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        raise ValueError("History: upload an .xlsx or .csv file.")
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if not raw:
        raise ValueError("History file is empty.")
    alias = {a: key for key, als in ALIASES.items() for a in als}
    cols = {i: alias[_norm(h)] for i, h in enumerate(raw[0]) if _norm(h) in alias}
    missing = [k for k in ("semester", "staff_id") if k not in cols.values()]
    if missing:
        raise ValueError("History file needs the columns: Semester, Staff ID (missing: "
                         + ", ".join(missing) + ").")
    rows = []
    for r in raw[1:]:
        row = {key: (r[i] if i < len(r) else None) for i, key in cols.items()}
        if row.get("semester") in (None, "") or row.get("staff_id") in (None, ""):
            continue
        row["semester"] = str(row["semester"]).strip()
        row["staff_id"] = str(row["staff_id"]).strip()
        rows.append(row)
    return rows
